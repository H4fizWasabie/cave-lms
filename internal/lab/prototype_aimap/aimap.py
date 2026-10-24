# PROTOTYPE, throwaway: answers "Can the AI map real spreadsheets accurately enough" (cave-lms #210).
# Never merged. Reads the private samples from $LAB_SAMPLES (default ~/cave-lab-samples); nothing from them is written
# into the repo. Usage:
#   python3 aimap.py preview raw|masked      print exactly what would be sent, per file
#   python3 aimap.py baseline                Go-rules-plus-looser-guessing stand-in, scored
#   python3 aimap.py ai raw|masked [runs]    ask OpenRouter (provider.zdr), score, report cost
# The AI call needs OPENROUTER_API_KEY; OPENROUTER_MODEL defaults to staging's openai/gpt-6-luna.
import datetime, glob, json, os, re, sys, urllib.request
import openpyxl

SAMPLES = os.path.expanduser(os.environ.get("LAB_SAMPLES", "~/cave-lab-samples"))
ROWS, COLS, CELL = 15, 40, 40  # grid sent per sheet
TEMPLATES = {  # required blanks first; Sales and Roster as decided in #215
    "Sales": (["date", "amount"], ["item", "category", "staff", "payment", "quantity"]),
    "Roster": (["staff", "date", "code"], ["department"]),
    "Finance": (["date", "amount"], ["category", "type", "description"]),
    "Inventory": (["item", "quantity"], ["reorder", "category", "price"]),
    "HR": (["name", "department"], ["joined", "status", "gender"]),
    "General": ([], []),
}
# ponytail: wider guess lists as 209 asks ("substring, wider lists"); the real lists get written with the Templates.
GUESS = {
    "date": ["tarikh", "date"], "amount": ["jumlah", "amaun", "amount", "total", "price", "harga"],
    "item": ["product name", "item", "barang", "product / service", "service name", "nama"],
    "quantity": ["kuantiti", "quantity", "qty", "actual stock", "stok", "baki", "closing balance", "balance"],
    "category": ["kategori", "category", "type"], "staff": ["staff", "pekerja", "created by", "doctor"],
    "payment": ["payment", "bayaran"], "code": ["code", "shift", "kod"], "department": ["jabatan", "department", "sheet"],
    "reorder": ["reorder", "minimum", "alert"], "price": ["selling price", "harga", "unit price", "price"],
    "name": ["nama", "name"], "type": ["jenis", "masuk/keluar"], "description": ["perkara", "butiran", "description"],
    "joined": ["join date", "start date", "tarikh mula"], "status": ["status"], "gender": ["jantina", "gender"],
}
NEEDS = {"date": "date", "amount": "number", "quantity": "number", "reorder": "number", "price": "number"}

# The right answer per file, agreed with Abah. Lists are accepted alternatives; None means not scored.
DAILY = [f"{d:02d}.09.26" for d in range(1, 31)]
SALES6 = [f"{m} ({k})" for m in ("June", "July", "August") for k in ("Service", "Product")]
ROSTER8 = ["Veterinary", "Surgery Dept", "Operation Room Supervisor", "Clinical Support Assistant", "Vet Nurse",
           "Customer Services", "Laboratory", "Pharmacy"]
ANSWERS = {
    "9. PUSVMC Daily": dict(template=["Sales"], sheets=DAILY, header=[8], unpivot=False,
                           req={"date": ["date", "sheet"], "amount": ["total sales", "total amount"]}),
    "Service & Product": dict(template=["Sales"], sheets=SALES6, header=[1], unpivot=False,
                              req={"date": ["created date & time"], "amount": ["price"]}),
    "High movement": dict(template=["General"], sheets=["all item available"], header=[1], unpivot=False, req={}),
    "roster": dict(template=["Roster"], sheets=ROSTER8, header=[6, 7], unpivot=(10, 2026),
                   req={"staff": ["b", "column b", "staff", "name", "staff name"], "date": ["date"],
                        "code": ["code", "shift", "shift code", "value"]}),
    "Test & Parameters": dict(template=["General"], sheets=None, sheet_any=["Sheet4", "Sheet5", "Sheet9", "ITEM name",
                              "Supplier Info"], header=[1], unpivot=False, req={}),
    "Stock_Balance": dict(template=["Inventory"], sheets=["Sheet1"], header=[1], unpivot=False,
                          req={"item": ["product name"], "quantity": ["actual stock"]}),
    "Stock_Movement": dict(template=["Inventory", "General"], sheets=["Sheet1"], header=[1], unpivot=False,
                           req={"item": ["product name"], "quantity": ["closing balance"]}),
    "Stock_Detail": dict(template=["General"], sheets=["Sheet1"], header=[1], unpivot=False, req={}),
}


def cell(v):
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.strftime("%d/%m/%Y %H:%M:%S") if (v.hour or v.minute) else v.strftime("%d/%m/%Y")
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip().replace("\n", " ")[:CELL]


def kind(s):
    s = s.strip()
    if s in ("", "-"):
        return "blank"
    if re.fullmatch(r"(RM\s*)?-?[\d,]*\.?\d+", s):
        return "number"
    if re.match(r"\d{1,4}[/.-]\d{1,2}[/.-]\d{2,4}", s):
        return "date"
    return "text"


def mask(s):  # letters to a/A, digits to 9, everything else kept: the shape without the content
    return "".join("9" if c.isdigit() else ("A" if c.isupper() else "a") if c.isalpha() else c for c in s)


def load(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = [[cell(v) for v in r[:COLS]] for r in ws.iter_rows(max_row=40, values_only=True)]
        out.append(dict(name=ws.title, nrows=ws.max_row, ncols=ws.max_column, rows=rows))
    return out


def header_row(rows):
    """209's rule, roughly: skip blank and title rows; the header is the first all-text row nearly as wide as the table."""
    width = max((sum(1 for c in r if c) for r in rows[:30]), default=0)
    for i, r in enumerate(rows[:30]):
        filled = [c for c in r if c]
        if filled and len(filled) >= 0.6 * width and all(kind(c) == "text" for c in filled):
            return i + 1
    return 1


def names(rows, h):
    r = rows[h - 1] if h - 1 < len(rows) else []
    return [c for c in r if c]


def day_grid(cols):
    days = {"mon", "tue", "wed", "thu", "fri", "sat", "sun", "isn", "sel", "rab", "kha", "jum", "sab", "ahd"}
    return sum(1 for c in cols if c.lower()[:3] in days or re.fullmatch(r"\d{1,2}", c)) >= 20


def payload(sheets, masked):
    """What goes to the model: every sheet's name and size, and a grid for the first 3 and last 2 sheets."""
    shown = sheets if len(sheets) <= 5 else sheets[:3] + sheets[-2:]
    parts = ["Sheets: " + "; ".join(f"{s['name']} ({s['nrows']} rows x {s['ncols']} cols)" for s in sheets)]
    for s in shown:
        h = header_row(s["rows"])
        lines = []
        for i, r in enumerate(s["rows"][:ROWS]):
            keep = not masked or i + 1 in (h, h - 1)  # masked: only Go's header guess and the row above stay as text
            cells = [c if keep else mask(c) for c in r]
            while cells and not cells[-1]:
                cells.pop()
            lines.append(f"{i + 1}: " + " | ".join(cells))
        parts.append(f"## Sheet {s['name']}\n" + "\n".join(lines))
    return "\n\n".join(parts)


def column_kind(s, h, col):
    vals = [kind(r[col]) for r in s["rows"][h:] if col < len(r) and kind(r[col]) != "blank"]
    return max(set(vals), key=vals.count) if vals else "blank"


def baseline(sheets):
    """Go rules plus looser guessing: first readable sheet, same-column sheets ticked, Template whose required blanks all match."""
    readable = [s for s in sheets if len(names(s["rows"], header_row(s["rows"]))) >= 3]
    s = readable[0] if readable else sheets[0]
    h = header_row(s["rows"])
    cols = names(s["rows"], h)
    norm = lambda cs: sorted(c.lower().strip() for c in cs)
    combine = [t["name"] for t in sheets if t is not s and norm(names(t["rows"], header_row(t["rows"]))) == norm(cols)]
    if day_grid(cols):
        cols = ["Staff", "Date", "Code"]  # the unpivot's own columns (215)
    if combine:
        cols = cols + (["Date"] if all(re.fullmatch(r"\d{1,2}[./-]\d{1,2}[./-]\d{2,4}", n) for n in [s["name"]] + combine) else ["Sheet"])
    header = s["rows"][h - 1] if h - 1 < len(s["rows"]) else []

    def guess(blank, used):
        for g in GUESS.get(blank, []):
            for c in cols:
                if c in used or g not in c.lower():
                    continue
                if blank in NEEDS and c in header and column_kind(s, h, header.index(c)) not in (NEEDS[blank], "blank"):
                    continue
                return c
        return ""

    for t in ["Sales", "Roster", "Inventory", "Finance", "HR"]:
        req, opt = TEMPLATES[t]
        used, mapping = set(), {}
        for b in req + opt:
            mapping[b] = guess(b, used)
            used.add(mapping[b])
        if all(mapping[b] for b in req):
            return dict(template=t, sheet=s["name"], combine=combine, header_row=h, two_row_header=False,
                        unpivot=day_grid(names(s["rows"], h)), month=0, year=0,
                        mapping=[{"blank": b, "column": c} for b, c in mapping.items()])
    return dict(template="General", sheet=s["name"], combine=combine, header_row=h, two_row_header=False,
                unpivot=day_grid(names(s["rows"], h)), month=0, year=0, mapping=[])


INSTRUCTIONS = """You plan how Cave LMS turns an uploaded spreadsheet into a dashboard. You see the workbook's sheets
(a grid of the first rows of some of them; cells are '|'-separated, empty cells are empty). Reply with a JSON plan
only; Cave's own code applies it and the Learner confirms it in a preview.
- template: Sales (what sells and when: blanks date, amount required; item, category, staff, payment, quantity
  optional), Roster (who works when: staff, date, code required; department optional), Finance (money in and out:
  date, amount required; category, type, description optional), Inventory (stock on hand: item, quantity required;
  reorder, category, price optional), HR (staff list: name, department required; joined, status, gender optional),
  or General (any table; no blanks) when no Template's question fits the data.
- sheet: the sheet holding the table. combine: other sheets with the same columns that belong in the same table
  (same columns after their own header row; not filtered copies of the sheet). Combining adds a column named Sheet
  holding each sheet's name, or named Date when every combined sheet name is a date.
- header_row: the 1-based row of the column names (title and blank rows above are dropped). two_row_header: true when
  the row above it is part of the header (group names over columns).
- unpivot: true when days run across columns (a staff x day grid of codes); Cave turns it into one row per
  staff per day with columns Staff, Date and Code (Staff is the grid's first text column). month, year: the grid's
  month and year (0 when not unpivoting).
- mapping: for each blank of the chosen template, the column name (or a column letter when the header cell is
  empty), or "" when no column fits. Never fill a blank with a column of the wrong meaning just to fill it."""

SCHEMA = {"type": "object", "additionalProperties": False,
          "required": ["template", "sheet", "combine", "header_row", "two_row_header", "unpivot", "month", "year", "mapping"],
          "properties": {
              "template": {"type": "string", "enum": list(TEMPLATES)}, "sheet": {"type": "string"},
              "combine": {"type": "array", "items": {"type": "string"}}, "header_row": {"type": "integer"},
              "two_row_header": {"type": "boolean"}, "unpivot": {"type": "boolean"},
              "month": {"type": "integer"}, "year": {"type": "integer"},
              "mapping": {"type": "array", "items": {"type": "object", "additionalProperties": False,
                          "required": ["blank", "column"],
                          "properties": {"blank": {"type": "string"}, "column": {"type": "string"}}}}}}


def ask(text):
    body = {"model": os.environ.get("OPENROUTER_MODEL", "openai/gpt-6-luna"), "max_tokens": 4000,
            "provider": {"zdr": True}, "usage": {"include": True},
            "messages": [{"role": "system", "content": INSTRUCTIONS}, {"role": "user", "content": text}],
            "response_format": {"type": "json_schema", "json_schema": {"name": "plan", "strict": True, "schema": SCHEMA}}}
    req = urllib.request.Request("https://openrouter.ai/api/v1/chat/completions", json.dumps(body).encode(),
                                 {"Authorization": "Bearer " + os.environ["OPENROUTER_API_KEY"],
                                  "Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=120) as r:
        reply = json.load(r)
    return json.loads(reply["choices"][0]["message"]["content"]), reply.get("usage", {}), reply.get("provider")


def score(key, plan):
    a = ANSWERS[key]
    got = {"template": plan["template"] in a["template"]}
    picked = [plan["sheet"]] + plan["combine"]
    if a["sheets"] is None:
        got["sheets"] = plan["sheet"] in a["sheet_any"] and not plan["combine"]
    else:
        got["sheets"] = sorted(set(picked)) == sorted(a["sheets"])
    got["header"] = plan["header_row"] in a["header"]
    got["unpivot"] = plan["unpivot"] == bool(a["unpivot"])  # month and year are always asked in the preview (215); reported, not scored
    m = {x["blank"]: x["column"].strip().lower() for x in plan["mapping"]}
    # a wrong Template is already counted once; General has no blanks
    got["mapping"] = not got["template"] or plan["template"] == "General" or \
        all(m.get(b, "") in ok for b, ok in a["req"].items())
    return got


def files():
    for path in sorted(glob.glob(os.path.join(SAMPLES, "*.xlsx"))):
        key = next(k for k in ANSWERS if k in os.path.basename(path))
        yield key, load(path)


def report(key, plan, got):
    ok = all(got.values())
    wrong = [k for k, v in got.items() if not v]
    print(f"{'RIGHT' if ok else 'wrong':5}  {key:20} {plan['template']:9} sheet={plan['sheet']!r} +{len(plan['combine'])} "
          f"header={plan['header_row']}{'(2)' if plan['two_row_header'] else ''} unpivot={plan['unpivot']}"
          f"{(plan['month'], plan['year']) if plan['unpivot'] else ''} "
          f"map={ {x['blank']: x['column'] for x in plan['mapping'] if x['column']} }  wrong={wrong}")
    return ok


if __name__ == "__main__":
    mode = sys.argv[1]
    if mode == "preview":
        for key, sheets in files():
            text = payload(sheets, sys.argv[2] == "masked")
            print(f"########## {key}  ({len(text)} chars)\n{text}\n")
    elif mode == "baseline":
        print(f"baseline: {sum(report(k, p, score(k, p)) for k, s in files() for p in [baseline(s)])}/8 right")
    elif mode == "ai":
        masked, runs = sys.argv[2] == "masked", int(sys.argv[3]) if len(sys.argv) > 3 else 1
        for run in range(runs):
            right, cost, toks = 0, 0.0, 0
            for key, sheets in files():
                plan, usage, provider = ask(payload(sheets, masked))
                right += report(key, plan, score(key, plan))
                cost += usage.get("cost", 0) or 0
                toks += usage.get("prompt_tokens", 0)
                print(f"       {provider} prompt={usage.get('prompt_tokens')} completion={usage.get('completion_tokens')} "
                      f"cost=${usage.get('cost')}")
            print(f"run {run + 1} ({'masked' if masked else 'raw'}): {right}/8 right, ${cost:.4f} for 8 uploads, "
                  f"{toks // 8} prompt tokens per upload\n")
